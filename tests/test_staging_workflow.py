"""
tests/test_staging_workflow.py
--------------------------------
The Gov't Data Upload's three-step workflow and the upload speed work:

  * an LGU Encoder uploads to the staging area -- nothing live changes;
  * an LGU Data Steward reviews: never their own upload, a rejection
    needs a reason;
  * a System Admin publishes approved data to the live tables (once).

Plus the fast .xlsx reader (same rows as openpyxl), the import lock not
being held by the follow-up work, retention of staged batches, and the
sidebar / audit labels that go with the workflow.
"""

import csv
import datetime
import os
import tempfile

import pytest

from app import create_app
from app.extensions import db
from app.ml.seed_data import BARANGAY_NAMES
from app.models import AuditLog, LguData, MarketData, Notification, SystemSetting, UploadBatch
from app.models.user import User

FOOD = "Food and Beverage"
RETAIL = "Wholesale and Retail Trade; Repair of Motor Vehicles and Motorcycles"
HEADER = ["permit_no", "barangay", "industry_type", "line_of_business", "business_name", "status"]
SAMPLE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "sample_data",
                      "MarketLens_LGU_Sample_Upload_Tarlac_City.xlsx")


@pytest.fixture
def app():
    app = create_app("testing")

    # The fixture keeps an app context open, and Flask reuses it for the
    # test client's requests -- so flask.g, where Flask-Login caches the
    # signed-in user, would carry one account into the next client's
    # requests. Start every request with no cached user, as a server does.
    def _fresh_login_cache():
        from flask import g

        g.pop("_login_user", None)

    app.before_request_funcs.setdefault(None, []).insert(0, _fresh_login_cache)
    with app.app_context():
        db.create_all()
        SystemSetting.ensure_defaults()
        for name, email, role, lgu_role in (
            ("Encoder One", "encoder@flow.test", "LGU", "encoder"),
            ("Steward One", "steward@flow.test", "LGU", "steward"),
            ("Steward Two", "steward2@flow.test", "LGU", "steward"),
            ("Admin One", "admin@flow.test", "Admin", None),
            ("Owner One", "sme@flow.test", "SME", None),
        ):
            user = User(name=name, email=email, role=role, lgu_role=lgu_role)
            user.set_password("password123")
            db.session.add(user)
        db.session.commit()
        yield app
        db.session.remove()
        db.drop_all()


def _uid(email):
    return User.query.filter_by(email=email).one().user_id


def _login(app, email):
    client = app.test_client()
    client.post("/login", data={"email": email, "password": "password123"})
    return client


def _register_csv(barangays=BARANGAY_NAMES[:3], per_barangay=4):
    handle = tempfile.NamedTemporaryFile("w", suffix=".csv", delete=False, encoding="utf-8", newline="")
    writer = csv.writer(handle)
    writer.writerow(HEADER)
    number = 0
    for barangay in barangays:
        for index in range(per_barangay):
            number += 1
            industry, line = ((FOOD, "Bakery"), (RETAIL, "Hardware"))[index % 2]
            writer.writerow([f"P-{number:05d}", barangay, industry, line, f"Shop {number}", "active"])
    handle.close()
    return handle.name


def _upload(client, path, dataset_type="MARKET_DATA", source="DTI"):
    with open(path, "rb") as handle:
        return client.post("/lgu/government-data-upload", data={
            "dataset_type": dataset_type, "source": source, "dataset_file": (handle, "permits.csv"),
        }, content_type="multipart/form-data", follow_redirects=True)


def _stage(app, uploader_email="encoder@flow.test"):
    from app.services.staging_service import stage_upload

    path = _register_csv()
    try:
        _count, status, message, batch_id = stage_upload("MARKET_DATA", path, "DTI", _uid(uploader_email),
                                                         "permits.csv")
    finally:
        os.unlink(path)
    assert status == "success", message
    return batch_id


# ---------------------------------------------------------------- step 1

def test_an_encoder_upload_goes_to_staging_and_stewards_are_told(app):
    path = _register_csv()
    try:
        page = _upload(_login(app, "encoder@flow.test"), path).get_data(as_text=True)
    finally:
        os.unlink(path)
    assert "Saved to the staging area as batch #" in page
    batch = UploadBatch.query.one()
    assert batch.status == "pending_review" and batch.barangays == 3 and batch.market_rows == 6
    assert MarketData.query.filter_by(source="DTI").count() == 0
    from app.services.data_import_service import has_active_lgu_data

    assert not has_active_lgu_data(), "city-wide recommendations stay off until the batch is published"
    told = {n.user_id for n in Notification.query.filter(Notification.message.contains("waiting for your review"))}
    assert told == {_uid("steward@flow.test"), _uid("steward2@flow.test")}


def test_a_data_steward_cannot_upload(app):
    path = _register_csv()
    try:
        page = _upload(_login(app, "steward@flow.test"), path).get_data(as_text=True)
    finally:
        os.unlink(path)
    assert "Gov't Data Review" in page.replace("&#39;", "'"), "a steward's data page is the review page"
    assert UploadBatch.query.count() == 0
    response = _login(app, "steward@flow.test").get("/lgu/government-data-upload")
    assert response.status_code == 302 and response.headers["Location"].endswith("/lgu/data-review")


def test_an_lgu_account_without_a_sub_role_is_an_encoder(app):
    user = User(name="Legacy LGU", email="legacy@flow.test", role="LGU")
    user.set_password("password123")
    db.session.add(user)
    db.session.commit()
    assert user.is_lgu_encoder() and not user.is_data_steward()
    assert user.lgu_role_label == "LGU Encoder"


# ---------------------------------------------------------------- step 2

def test_only_data_stewards_open_data_review(app):
    assert _login(app, "steward@flow.test").get("/lgu/data-review").status_code == 200
    for email in ("encoder@flow.test", "admin@flow.test", "sme@flow.test"):
        assert _login(app, email).get("/lgu/data-review").status_code == 403, email


def test_a_steward_approves_and_the_admins_are_told(app):
    batch_id = _stage(app)
    client = _login(app, "steward@flow.test")
    review = client.get(f"/lgu/data-review/{batch_id}", follow_redirects=True).get_data(as_text=True)
    assert "Barangay rows" in review and BARANGAY_NAMES[0] in review and 'id="reviewForm"' in review

    page = client.post(f"/lgu/data-review/{batch_id}", data={"decision": "approve", "note": "Figures match BPLO."},
                       follow_redirects=True).get_data(as_text=True)
    assert f"Batch #{batch_id} approved" in page
    batch = db.session.get(UploadBatch, batch_id)
    db.session.refresh(batch)
    assert batch.status == "approved" and batch.reviewed_by == _uid("steward@flow.test")
    assert batch.review_note == "Figures match BPLO." and batch.payload
    assert Notification.query.filter_by(user_id=_uid("admin@flow.test")).filter(
        Notification.message.contains("ready to publish")).count() == 1
    assert AuditLog.query.filter_by(action="staging_review_approve").count() == 1
    assert MarketData.query.filter_by(source="DTI").count() == 0, "approval alone publishes nothing"


def test_a_steward_never_reviews_their_own_upload(app):
    from app.services.staging_service import review

    steward = User.query.filter_by(email="steward@flow.test").one()
    steward_two = User.query.filter_by(email="steward2@flow.test").one()
    steward.lgu_role = "encoder"
    db.session.commit()
    batch_id = _stage(app, "steward@flow.test")
    steward.lgu_role = "steward"
    db.session.commit()

    ok, message = review(batch_id, steward, approve=True)
    assert not ok and "another Data Steward" in message
    page = _login(app, "steward@flow.test").get(f"/lgu/data-review/{batch_id}").get_data(as_text=True)
    assert 'id="reviewForm"' not in page
    ok, _message = review(batch_id, steward_two, approve=True)
    assert ok


def test_rejecting_needs_a_reason_and_clears_the_staged_rows(app):
    batch_id = _stage(app)
    client = _login(app, "steward@flow.test")
    page = client.post(f"/lgu/data-review/{batch_id}", data={"decision": "reject", "note": "  "},
                       follow_redirects=True).get_data(as_text=True)
    assert "give the reason for rejecting" in page
    assert db.session.get(UploadBatch, batch_id).status == "pending_review"

    client.post(f"/lgu/data-review/{batch_id}", data={"decision": "reject", "note": "Wrong barangay column."})
    db.session.expire_all()
    batch = db.session.get(UploadBatch, batch_id)
    assert batch.status == "rejected" and batch.payload is None
    encoder_page = _login(app, "encoder@flow.test").get("/lgu/government-data-upload").get_data(as_text=True)
    assert "Wrong barangay column." in encoder_page and "Rejected" in encoder_page


# ---------------------------------------------------------------- step 3

def _approved(app):
    from app.services.staging_service import review

    batch_id = _stage(app)
    ok, message = review(batch_id, User.query.filter_by(email="steward@flow.test").one(), approve=True)
    assert ok, message
    return batch_id


def test_only_an_admin_publishes(app):
    batch_id = _approved(app)
    for email in ("encoder@flow.test", "steward@flow.test", "sme@flow.test"):
        client = _login(app, email)
        assert client.get("/admin/publish-data").status_code == 403, email
        assert client.get("/admin/datasets").status_code == 403, email
        assert client.post(f"/admin/publish-data/{batch_id}/publish").status_code == 403, email
    assert db.session.get(UploadBatch, batch_id).status == "approved"


def test_publishing_makes_the_batch_live_once(app):
    from app.services.data_import_service import has_active_lgu_data

    batch_id = _approved(app)
    client = _login(app, "admin@flow.test")
    listing = client.get("/admin/publish-data", follow_redirects=True).get_data(as_text=True)
    assert f"#{batch_id}" in listing and "Publish approved government data" in listing, "the panel is on Datasets"

    page = client.post(f"/admin/publish-data/{batch_id}/publish", follow_redirects=True).get_data(as_text=True)
    assert f"Batch #{batch_id} published" in page
    db.session.expire_all()
    batch = db.session.get(UploadBatch, batch_id)
    assert batch.status == "published" and batch.payload is None
    assert batch.published_by == _uid("admin@flow.test")
    assert MarketData.query.filter_by(source="DTI").count() == 6
    assert LguData.query.filter(LguData.barangay.in_(BARANGAY_NAMES[:3])).count() == 3
    assert all(row.upload_date == datetime.date.today() for row in LguData.query.all())
    assert has_active_lgu_data()
    assert AuditLog.query.filter_by(action="staging_publish").count() == 1
    assert Notification.query.filter_by(user_id=_uid("encoder@flow.test")).filter(
        Notification.message.contains("is now live")).count() == 1

    again = client.post(f"/admin/publish-data/{batch_id}/publish", follow_redirects=True).get_data(as_text=True)
    assert "not approved for publishing" in again
    assert MarketData.query.filter_by(source="DTI").count() == 6, "a second click writes nothing"


def test_a_pending_batch_cannot_be_published(app):
    from app.services.staging_service import publish

    batch_id = _stage(app)
    ok, _message = publish(batch_id, User.query.filter_by(email="admin@flow.test").one())
    assert not ok
    assert MarketData.query.filter_by(source="DTI").count() == 0


def test_an_admin_can_reject_an_approved_batch_with_a_reason(app):
    batch_id = _approved(app)
    client = _login(app, "admin@flow.test")
    client.post(f"/admin/publish-data/{batch_id}/reject", data={"reason": ""})
    assert db.session.get(UploadBatch, batch_id).status == "approved"
    client.post(f"/admin/publish-data/{batch_id}/reject", data={"reason": "Superseded by a newer register."})
    db.session.expire_all()
    batch = db.session.get(UploadBatch, batch_id)
    assert batch.status == "rejected" and batch.payload is None
    assert "Superseded by a newer register." in batch.review_note and "Admin One" in batch.review_note
    assert AuditLog.query.filter_by(action="staging_admin_reject").count() == 1


# ------------------------------------------------------------- accounts

def test_an_admin_sets_the_lgu_role(app):
    client = _login(app, "admin@flow.test")
    encoder = _uid("encoder@flow.test")
    assert "LGU Data Steward" in client.get("/admin/users").get_data(as_text=True)
    client.post(f"/admin/users/{encoder}/lgu-role", data={"lgu_role": "steward"})
    db.session.expire_all()
    assert db.session.get(User, encoder).is_data_steward()
    assert AuditLog.query.filter_by(action="admin_set_lgu_role").count() == 1
    client.post(f"/admin/users/{_uid('sme@flow.test')}/lgu-role", data={"lgu_role": "steward"})
    assert db.session.get(User, _uid("sme@flow.test")).lgu_role is None, "only LGU accounts have an LGU role"


def test_an_admin_creates_a_data_steward(app):
    from tests.account_fixtures import STRONG_PASSWORD

    _login(app, "admin@flow.test").post("/admin/users/create", data={
        "first_name": "Rosa", "last_name": "Cruz", "email": "rosa@flow.test", "password": STRONG_PASSWORD,
        "role": "LGU_STEWARD",
    })
    user = User.query.filter_by(email="rosa@flow.test").one()
    assert user.role == "LGU" and user.is_data_steward()


def test_the_sidebar_shows_each_step_to_the_right_people(app):
    _approved(app)
    _stage(app)
    steward = _login(app, "steward@flow.test").get("/lgu/data-review").get_data(as_text=True).replace("&#39;", "'")
    assert "Gov't Data Review" in steward and "Uploads waiting for your review" in steward
    assert 'href="/lgu/government-data-upload"' not in steward, "one data page for a steward, not two"
    encoder = _login(app, "encoder@flow.test").get("/lgu/government-data-upload").get_data(as_text=True)
    assert 'href="/lgu/data-review"' not in encoder and "Gov't Data Upload" in encoder.replace("&#39;", "'")
    admin = _login(app, "admin@flow.test").get("/admin/datasets").get_data(as_text=True)
    assert 'href="/admin/publish-data"' not in admin and "Approved batches waiting to be published" in admin
    assert admin.index('id="publish-data"') < admin.index("LGU Data (permits"), "the publish panel is above the tables"


def test_every_workflow_action_has_an_audit_label():
    from app.utils.audit_labels import category_for, importance_for, label_for

    for action in ("staging_review_approve", "staging_review_reject", "staging_publish", "staging_admin_reject",
                   "admin_set_lgu_role"):
        assert label_for(action) != action
        assert category_for(action) != "other"
        assert importance_for(action) == "high"


# ------------------------------------------------------------ retention

def test_staged_batches_follow_retention(app):
    from app.services.retention_service import STAGING_EXPIRY_DAYS, apply_retention

    stale_id = _stage(app)
    old_id = _approved(app)
    stale = db.session.get(UploadBatch, stale_id)
    stale.uploaded_at = datetime.datetime.utcnow() - datetime.timedelta(days=STAGING_EXPIRY_DAYS + 1)
    old = db.session.get(UploadBatch, old_id)
    old.uploaded_at = datetime.datetime.utcnow() - datetime.timedelta(days=4000)
    db.session.commit()

    counts = apply_retention()
    assert counts["upload_batches"] == 1 and counts["staged_expired"] == 1
    db.session.expire_all()
    assert db.session.get(UploadBatch, old_id) is None
    stale = db.session.get(UploadBatch, stale_id)
    assert stale.status == "rejected" and stale.payload is None and "Expired" in stale.review_note


# ----------------------------------------------------------------- speed

def test_the_import_lock_is_free_while_follow_up_work_runs(app, monkeypatch):
    from app.services import data_import_service, upload_jobs

    seen = []
    monkeypatch.setattr(data_import_service, "_run_alert_sweeps", lambda: seen.append(upload_jobs._RUN_LOCK.locked()))
    monkeypatch.setattr(upload_jobs, "warm_after_upload", lambda superseded=None: seen.append("warm"))
    app.config.update(UPLOAD_IN_BACKGROUND=False, WARM_AFTER_UPLOAD=True)

    path = _register_csv()
    job = upload_jobs.start_upload(app, _uid("encoder@flow.test"), path, "permits.csv", "MARKET_DATA", "DTI")
    assert job.status == "success" and job.batch_id
    assert seen == [], "staging an upload runs no alert sweep and no warm-up"
    assert not upload_jobs._RUN_LOCK.locked()

    upload_jobs.schedule_followups(app)
    assert seen == [False, "warm"], "publishing's follow-up work never holds the import lock"


def _write_workbook(path):
    import openpyxl

    workbook = openpyxl.Workbook()
    notes = workbook.active
    notes.title = "Read me"
    notes.append(["How to fill in this workbook"])
    data = workbook.create_sheet("Permits")
    data.append(["permit_no", "barangay", "industry_type", "status", "capital", "effective_date", "open"])
    data.append(["P-1", "San Roque", FOOD, "active", 15000, datetime.datetime(2026, 3, 1), True])
    data.append(["P-2", "Mabini", RETAIL, None, 2500.75, datetime.date(2025, 12, 31), False])
    data.append(["P-3", "Ligtasan", "Café & bakery <ñ>", "expired", None, None, None])
    data["D4"] = "  spaced  "
    workbook.save(path)


def _openpyxl_rows(path, sheet_index):
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return [list(row) for row in workbook.worksheets[sheet_index].iter_rows(values_only=True)]
    finally:
        workbook.close()


def test_the_fast_xlsx_reader_reads_what_openpyxl_reads(tmp_path):
    from app.utils.xlsx_reader import read_data_sheet

    path = str(tmp_path / "permits.xlsx")
    _write_workbook(path)
    header, rows = read_data_sheet(path, lambda cells: "barangay" in cells)
    expected = _openpyxl_rows(path, 1)
    assert [list(header)] + [list(row) for row in rows] == expected


@pytest.mark.skipif(not os.path.exists(SAMPLE), reason="sample workbook not bundled")
def test_the_fast_xlsx_reader_matches_openpyxl_on_the_sample():
    import openpyxl

    from app.utils.xlsx_reader import read_data_sheet

    header, rows = read_data_sheet(SAMPLE, lambda cells: "barangay" in {str(c).strip().lower() for c in cells})
    workbook = openpyxl.load_workbook(SAMPLE, read_only=True, data_only=True)
    try:
        index = next(i for i, sheet in enumerate(workbook.worksheets)
                     if "barangay" in {str(c).strip().lower() for c in next(sheet.iter_rows(max_row=1, values_only=True))})
    finally:
        workbook.close()
    expected = _openpyxl_rows(SAMPLE, index)
    assert [list(header)] + [list(row) for row in rows] == expected


# ------------------------------------------- corrections and archiving

def _request_changes(app, batch_id, note="Please fix the zoning of the first barangay."):
    _login(app, "steward@flow.test").post(f"/lgu/data-review/{batch_id}", data={"decision": "changes", "note": note})
    db.session.expire_all()
    return db.session.get(UploadBatch, batch_id)


def test_a_steward_can_send_a_batch_back_with_a_message(app):
    batch_id = _stage(app)
    client = _login(app, "steward@flow.test")
    page = client.post(f"/lgu/data-review/{batch_id}", data={"decision": "changes", "note": ""},
                       follow_redirects=True).get_data(as_text=True)
    assert "what the encoder should fix" in page
    batch = _request_changes(app, batch_id)
    assert batch.status == "changes_requested" and batch.payload, "the rows are kept for the encoder to fix"
    assert batch.last_message["kind"] == "changes_requested" and "zoning" in batch.last_message["text"]
    assert Notification.query.filter_by(user_id=_uid("encoder@flow.test")).filter(
        Notification.message.contains("needs changes")).count() == 1
    assert AuditLog.query.filter_by(action="staging_request_changes").count() == 1
    upload_page = _login(app, "encoder@flow.test").get("/lgu/government-data-upload").get_data(as_text=True)
    assert "asked for changes" in upload_page and "zoning of the first barangay" in upload_page
    assert f'href="/lgu/uploads/{batch_id}/edit"' in upload_page


def test_the_encoder_edits_the_rows_and_it_goes_back_for_review(app):
    from app.services.staging_service import staged_rows

    batch_id = _stage(app)
    _request_changes(app, batch_id)
    client = _login(app, "encoder@flow.test")
    edit_page = client.get(f"/lgu/uploads/{batch_id}/edit").get_data(as_text=True)
    assert "Edit barangay rows" in edit_page and "zoning of the first barangay" in edit_page
    rows = staged_rows(db.session.get(UploadBatch, batch_id))
    form = {"barangay": [r["barangay"] for r in rows], "note": "Fixed the zoning."}
    for i, r in enumerate(rows):
        form[f"zoning_info-{i}"] = "Commercial" if i == 0 else (r.get("zoning_info") or "Residential")
        form[f"permit_count-{i}"] = "-3" if i == 1 else str(r.get("permit_count") or "")
    bad = client.post(f"/lgu/uploads/{batch_id}/edit", data=form, follow_redirects=True).get_data(as_text=True)
    assert "cannot be negative" in bad
    form[f"permit_count-1"] = "7"
    client.post(f"/lgu/uploads/{batch_id}/edit", data=form)
    db.session.expire_all()
    batch = db.session.get(UploadBatch, batch_id)
    rows = staged_rows(batch)
    assert batch.status == "pending_review" and rows[0]["zoning_info"] == "Commercial" and rows[1]["permit_count"] == 7
    assert batch.last_message["kind"] == "edited" and batch.last_message["text"] == "Fixed the zoning."
    assert AuditLog.query.filter_by(action="staging_edit_rows").count() == 1
    from app.services.staging_service import review

    ok, _ = review(batch_id, User.query.filter_by(email="steward@flow.test").one(), "approve")
    assert ok, "the corrected batch can be approved"


def test_the_encoder_can_replace_the_file(app):
    batch_id = _stage(app)
    path = _register_csv(barangays=BARANGAY_NAMES[:2], per_barangay=2)
    try:
        with open(path, "rb") as handle:
            _login(app, "encoder@flow.test").post(f"/lgu/uploads/{batch_id}/edit", data={
                "action": "replace", "note": "Cleaned list.", "dataset_file": (handle, "cleaned.csv"),
            }, content_type="multipart/form-data")
    finally:
        os.unlink(path)
    db.session.expire_all()
    batch = db.session.get(UploadBatch, batch_id)
    assert batch.filename == "cleaned.csv" and batch.barangays == 2 and batch.status == "pending_review"


def test_only_the_uploader_edits_and_only_before_approval(app):
    batch_id = _stage(app)
    other = User(name="Encoder Two", email="encoder2@flow.test", role="LGU", lgu_role="encoder")
    other.set_password("password123")
    db.session.add(other)
    db.session.commit()
    assert _login(app, "encoder2@flow.test").get(f"/lgu/uploads/{batch_id}/edit").status_code == 403
    from app.services.staging_service import edit_rows, review

    review(batch_id, User.query.filter_by(email="steward@flow.test").one(), "approve")
    ok, message = edit_rows(batch_id, User.query.filter_by(email="encoder@flow.test").one(), {}, "late fix")
    assert not ok and "can no longer be changed" in message


def test_the_encoder_can_archive_before_publishing(app):
    batch_id = _approved(app)
    _login(app, "encoder@flow.test").post(f"/lgu/uploads/{batch_id}/archive")
    db.session.expire_all()
    batch = db.session.get(UploadBatch, batch_id)
    assert batch.status == "archived" and batch.payload is None
    from app.services.staging_service import publish

    ok, _ = publish(batch_id, User.query.filter_by(email="admin@flow.test").one())
    assert not ok and MarketData.query.filter_by(source="DTI").count() == 0, "an archived batch never goes live"
    assert AuditLog.query.filter_by(action="staging_archive").count() == 1


def test_the_review_page_lists_recently_published_data(app):
    batch_id = _approved(app)
    from app.services.staging_service import publish

    publish(batch_id, User.query.filter_by(email="admin@flow.test").one())
    page = _login(app, "steward@flow.test").get("/lgu/data-review").get_data(as_text=True)
    assert "Recently published by the System Admin" in page and f"#{batch_id}" in page and "Admin One" in page


# ---------------------------------------------------------------- scope

def test_a_barangay_encoder_stages_only_their_barangay(app):
    from app.services.staging_service import stage_upload

    user = User(name="Brgy Encoder", email="brgy@flow.test", role="LGU", lgu_role="encoder",
                lgu_scope="barangay", lgu_barangay=BARANGAY_NAMES[0])
    user.set_password("password123")
    db.session.add(user)
    db.session.commit()
    path = _register_csv(barangays=BARANGAY_NAMES[:3])
    try:
        _count, status, message, batch_id = stage_upload("MARKET_DATA", path, "DTI", user.user_id, "all.csv")
    finally:
        os.unlink(path)
    assert status == "failed" and "covers Barangay" in message and batch_id is None
    path = _register_csv(barangays=BARANGAY_NAMES[:1])
    try:
        _count, status, message, batch_id = stage_upload("MARKET_DATA", path, "DTI", user.user_id, "own.csv")
    finally:
        os.unlink(path)
    assert status == "success" and db.session.get(UploadBatch, batch_id).scope_barangay == BARANGAY_NAMES[0]


def test_a_barangay_steward_reviews_only_their_barangay(app):
    from app.services.staging_service import review

    city_batch = _stage(app)
    steward = User(name="Brgy Steward", email="brgysteward@flow.test", role="LGU", lgu_role="steward",
                   lgu_scope="barangay", lgu_barangay=BARANGAY_NAMES[5])
    steward.set_password("password123")
    db.session.add(steward)
    db.session.commit()
    page = _login(app, "brgysteward@flow.test").get("/lgu/data-review").get_data(as_text=True)
    assert "Nothing is waiting for review" in page
    ok, message = review(city_batch, steward, "approve")
    assert not ok and "outside your scope" in message
